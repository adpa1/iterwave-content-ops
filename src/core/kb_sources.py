"""选题库来源：本地路径 / 页面上传 / 远程地址。

与部署位置无关：
- 相对路径一律相对项目根目录解析（不依赖启动目录、不假设 macOS）
- 上传文件存到 <data_dir>/kb/uploads/
- 远程地址同步到 <data_dir>/kb/remote_cache.<ext>，KbReader 读缓存

对外展示路径时一律走 display_path()：项目内显示相对路径，项目外只显示文件名，
不把服务器目录结构（用户名、私有目录）暴露到页面或接口。
"""

from __future__ import annotations

import csv
import io
import json
import logging
import os
import re
import xml.etree.ElementTree as ET
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any, Optional
from urllib.parse import urlsplit

from .safe_http import SafeFetchError, fetch_public_bytes

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent.parent.parent

UPLOAD_EXTS = {".csv", ".jsonl", ".ndjson", ".json", ".db", ".sqlite", ".sqlite3", ".xlsx"}
MAX_UPLOAD_BYTES = 20 * 1024 * 1024
MAX_REMOTE_BYTES = 20 * 1024 * 1024
MAX_LOCAL_SOURCE_BYTES = 100 * 1024 * 1024
MAX_TOPIC_ROWS = 100_000
MAX_XLSX_EXPANDED_BYTES = 40 * 1024 * 1024
MAX_XLSX_ENTRIES = 1000
REMOTE_TIMEOUT = 30


class KbSourceError(ValueError):
    """选题库来源错误（消息可直接展示给用户，不含服务器路径）。"""


# ---------------------------------------------------------------------------
# 路径
# ---------------------------------------------------------------------------

def data_dir(cfg) -> Path:
    d = Path(cfg.paths.data_dir)
    return d if d.is_absolute() else PROJECT_ROOT / d


def kb_dir(cfg) -> Path:
    d = data_dir(cfg) / "kb"
    d.mkdir(parents=True, exist_ok=True)
    try:
        d.chmod(0o700)
    except OSError:
        logger.warning("无法收紧选题库目录权限：%s", d.name)
    return d


def remote_cache_path(cfg) -> Path:
    fmt = _remote_format(cfg.topic_kb)
    ext = {"csv": ".csv", "json": ".json", "jsonl": ".jsonl", "rss": ".jsonl"}.get(fmt, ".json")
    return kb_dir(cfg) / f"remote_cache{ext}"


def resolve_kb_path(cfg) -> Optional[Path]:
    """当前生效的选题库文件路径；未配置返回 None。"""
    kb = cfg.topic_kb
    st = getattr(kb, "source_type", "local") or "local"
    if st == "remote":
        if not (kb.remote_url or "").strip():
            return None
        return current_remote_cache(cfg) or remote_cache_path(cfg)
    raw = (kb.path or "").strip()
    if not raw:
        return None
    p = Path(os.path.expanduser(raw))
    return p if p.is_absolute() else PROJECT_ROOT / p


def display_path(p: Optional[Path | str]) -> str:
    """页面展示用路径：项目内 → 相对路径；项目外 → 只显示文件名。"""
    if not p:
        return ""
    p = Path(p)
    try:
        return str(p.resolve().relative_to(PROJECT_ROOT.resolve()))
    except (ValueError, OSError):
        return f"…/{p.name}"


_ABS_PATH_RE = re.compile(r"(/(?:Users|home|root|var|opt|srv|mnt|data|tmp)/[^\s:'\"，。]+)|([A-Za-z]:\\[^\s:'\"，。]+)")


def scrub_paths(text: str) -> str:
    """把错误信息里的服务器绝对路径替换为文件名。"""
    if not text:
        return text
    return _ABS_PATH_RE.sub(lambda m: display_path(m.group(0)), str(text))


def source_summary(cfg) -> dict[str, Any]:
    """供页面展示的选题库来源摘要（不含绝对路径、不含密钥）。"""
    kb = cfg.topic_kb
    st = getattr(kb, "source_type", "local") or "local"
    p = resolve_kb_path(cfg)
    out: dict[str, Any] = {
        "source_type": st,
        "configured": p is not None,
        "display_path": display_path(p) if p else "",
        "exists": bool(p and p.exists()),
        "remote_url": _mask_url(kb.remote_url) if st == "remote" else "",
        "remote_format": kb.remote_format,
        "remote_header": kb.remote_header,
        "remote_token_set": bool(kb.remote_token_env and os.environ.get(kb.remote_token_env)),
        "default_dimension": kb.default_dimension,
    }
    if p and p.exists():
        st_ = p.stat()
        out["size_bytes"] = st_.st_size
        out["updated_at"] = datetime.fromtimestamp(st_.st_mtime).isoformat(timespec="seconds")
    return out


def _mask_url(url: str) -> str:
    """用于界面展示的远程地址只保留协议和主机，路径及查询部分可能包含凭据。"""
    if not url:
        return ""
    try:
        parsed = urlsplit(url)
        if parsed.scheme.lower() not in ("http", "https") or not parsed.hostname:
            return "远程地址"
        host = parsed.hostname.lower()
        if ":" in host and not host.startswith("["):
            host = f"[{host}]"
        port = parsed.port
        default_port = 443 if parsed.scheme.lower() == "https" else 80
        authority = host if port is None or port == default_port else f"{host}:{port}"
        return f"{parsed.scheme.lower()}://{authority}/…"
    except ValueError:
        return "远程地址"


# ---------------------------------------------------------------------------
# 上传
# ---------------------------------------------------------------------------

def save_upload(cfg, filename: str, content: bytes) -> Path:
    name = Path(filename or "").name
    ext = Path(name).suffix.lower()
    if ext not in UPLOAD_EXTS:
        raise KbSourceError(
            f"不支持的文件类型 {ext or '（无后缀）'}，支持 CSV、Excel(.xlsx)、JSON、JSONL、SQLite(.db)"
        )
    if not content:
        raise KbSourceError("文件是空的")
    if len(content) > MAX_UPLOAD_BYTES:
        raise KbSourceError("文件超过 20MB，请拆分后上传，或改用远程地址")
    safe = re.sub(r"[^\w.\-\u4e00-\u9fff]+", "_", Path(name).stem)[:60] or "topics"
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    dest_dir = kb_dir(cfg) / "uploads"
    dest_dir.mkdir(parents=True, exist_ok=True)
    try:
        dest_dir.chmod(0o700)
    except OSError:
        logger.warning("无法收紧上传目录权限：%s", dest_dir.name)
    dest = dest_dir / f"{safe}_{stamp}{ext}"
    dest.write_bytes(content)
    try:
        dest.chmod(0o600)
    except OSError:
        logger.warning("无法收紧上传文件权限：%s", dest.name)
    logger.info("📥 选题库上传: %s (%d bytes)", dest.name, len(content))
    return dest


def project_relative(p: Path) -> str:
    try:
        return str(p.resolve().relative_to(PROJECT_ROOT.resolve()))
    except ValueError:
        return str(p)


# ---------------------------------------------------------------------------
# 远程同步
# ---------------------------------------------------------------------------

def _remote_format(kb) -> str:
    fmt = (kb.remote_format or "auto").lower()
    if fmt != "auto":
        return fmt
    url = (kb.remote_url or "").lower().split("?")[0]
    for ext, f in ((".csv", "csv"), (".jsonl", "jsonl"), (".ndjson", "jsonl"),
                   (".json", "json"), (".xml", "rss"), (".rss", "rss"), ("/feed", "rss"),
                   ("/rss", "rss"), (".atom", "rss")):
        if url.endswith(ext):
            return f
    return "auto"


def _sniff_format(body: bytes, content_type: str) -> str:
    ct = (content_type or "").lower()
    head = body.lstrip()[:200].lower()
    if "xml" in ct or "rss" in ct or head.startswith(b"<?xml") or b"<rss" in head or b"<feed" in head:
        return "rss"
    if "csv" in ct:
        return "csv"
    if head.startswith(b"[") or head.startswith(b"{"):
        try:
            json.loads(body.decode("utf-8"))
            return "json"
        except Exception:
            return "jsonl"
    return "csv"


def sync_remote(cfg) -> dict[str, Any]:
    """下载远程选题库到本地缓存。只允许公网 HTTP(S)，流式限制响应大小。"""
    kb = cfg.topic_kb
    url = (kb.remote_url or "").strip()
    if not url:
        raise KbSourceError("还没有填写远程地址")
    headers = {"User-Agent": "IterWave/1.0"}
    if kb.remote_header and kb.remote_token_env:
        token = os.environ.get(kb.remote_token_env)
        if token:
            headers[kb.remote_header] = token
    try:
        body, content_type, _final_url = fetch_public_bytes(
            url, headers=headers, timeout=REMOTE_TIMEOUT, max_bytes=MAX_REMOTE_BYTES
        )
    except SafeFetchError as exc:
        raise KbSourceError(str(exc)) from exc

    fmt = _remote_format(kb)
    if fmt == "auto":
        fmt = _sniff_format(body, content_type)

    rows = 0
    if fmt == "rss":
        items = parse_feed(body, default_dimension=kb.default_dimension, source=_mask_url(url))
        if len(items) > MAX_TOPIC_ROWS:
            raise KbSourceError(f"远程选题库超过 {MAX_TOPIC_ROWS:,} 行上限")
        dest = kb_dir(cfg) / "remote_cache.jsonl"
        with open(dest, "w", encoding="utf-8") as f:
            for it in items:
                f.write(json.dumps(it, ensure_ascii=False) + "\n")
        rows = len(items)
    else:
        if fmt not in ("csv", "json", "jsonl"):
            raise KbSourceError("远程文件格式不支持")
        ext = {"csv": ".csv", "json": ".json", "jsonl": ".jsonl"}[fmt]
        dest = kb_dir(cfg) / f"remote_cache{ext}"
        dest.write_bytes(body)
        try:
            dest.chmod(0o600)
        except OSError:
            logger.warning("无法收紧远程缓存文件权限：%s", dest.name)
        rows = _count_rows(dest, fmt)
        if rows > MAX_TOPIC_ROWS:
            dest.unlink(missing_ok=True)
            raise KbSourceError(f"远程选题库超过 {MAX_TOPIC_ROWS:,} 行上限")
    # 清掉其它格式的旧缓存，避免读到过期文件
    for old in kb_dir(cfg).glob("remote_cache.*"):
        if old != dest:
            old.unlink(missing_ok=True)
    # remote_format=auto 时记住探测结果，读取时按真实格式解析
    _write_meta(cfg, {"format": fmt, "synced_at": datetime.now().isoformat(timespec="seconds"),
                      "rows": rows})
    logger.info("🌐 远程选题库同步完成: %s 条（%s）", rows, fmt)
    return {"ok": True, "format": fmt, "rows": rows,
            "synced_at": datetime.now().isoformat(timespec="seconds")}

def _meta_path(cfg) -> Path:
    return kb_dir(cfg) / "remote_meta.json"


def _write_meta(cfg, meta: dict) -> None:
    path = _meta_path(cfg)
    path.write_text(json.dumps(meta, ensure_ascii=False), encoding="utf-8")
    try:
        path.chmod(0o600)
    except OSError:
        logger.warning("无法收紧远程选题库元数据权限：%s", path.name)


def remote_meta(cfg) -> dict:
    p = _meta_path(cfg)
    if not p.exists():
        return {}
    try:
        return json.loads(p.read_text(encoding="utf-8"))
    except Exception:
        return {}


def current_remote_cache(cfg) -> Optional[Path]:
    caches = sorted(kb_dir(cfg).glob("remote_cache.*"))
    caches = [c for c in caches if c.suffix in (".csv", ".json", ".jsonl")]
    return caches[0] if caches else None


def _count_rows(p: Path, fmt: str) -> int:
    try:
        if fmt == "csv":
            with open(p, encoding="utf-8-sig") as f:
                return max(0, sum(1 for _ in csv.reader(f)) - 1)
        if fmt == "jsonl":
            with open(p, encoding="utf-8") as f:
                return sum(1 for line in f if line.strip())
        data = json.loads(p.read_text(encoding="utf-8"))
        return len(extract_json_rows(data))
    except Exception:
        return 0


def extract_json_rows(data: Any) -> list[dict]:
    """JSON 里找出行列表：数组本身，或 {items|data|list|records|rows: [...]}。"""
    if isinstance(data, list):
        return [r for r in data if isinstance(r, dict)]
    if isinstance(data, dict):
        for key in ("items", "data", "list", "records", "rows", "topics", "results"):
            v = data.get(key)
            if isinstance(v, list):
                return [r for r in v if isinstance(r, dict)]
            if isinstance(v, dict):
                inner = extract_json_rows(v)
                if inner:
                    return inner
    return []


# ---------------------------------------------------------------------------
# RSS / Atom
# ---------------------------------------------------------------------------

def _strip_tags(s: str) -> str:
    return re.sub(r"<[^>]+>", "", s or "").strip()


def parse_feed(body: bytes | str, *, default_dimension: str = "综合", source: str = "") -> list[dict]:
    """把 RSS 2.0 / Atom 解析成选题行（通用字段名）。"""
    if isinstance(body, bytes):
        text = body.decode("utf-8", errors="replace")
    else:
        text = body
    try:
        root = ET.fromstring(text)
    except ET.ParseError as exc:
        raise KbSourceError("RSS 内容解析失败，请确认地址是 RSS/Atom 订阅") from exc

    def local(tag: str) -> str:
        return tag.split("}", 1)[-1]

    out: list[dict] = []
    feed_title = ""
    for el in root.iter():
        if local(el.tag) == "title" and not feed_title:
            feed_title = (el.text or "").strip()
            break
    for el in root.iter():
        name = local(el.tag)
        if name not in ("item", "entry"):
            continue
        row: dict[str, str] = {}
        cats: list[str] = []
        for child in el:
            cn = local(child.tag)
            if cn == "title":
                row["title"] = _strip_tags(child.text or "")
            elif cn in ("description", "summary", "content") and "summary" not in row:
                row["summary"] = _strip_tags(child.text or "")[:500]
            elif cn == "link":
                row["link"] = (child.text or "").strip() or child.attrib.get("href", "")
            elif cn in ("pubDate", "published", "updated", "date") and "date" not in row:
                row["date"] = (child.text or "").strip()
            elif cn == "category":
                c = (child.text or child.attrib.get("term", "")).strip()
                if c:
                    cats.append(c)
        if not row.get("title"):
            continue
        row["dimension"] = cats[0] if cats else default_dimension
        row["source"] = feed_title or source
        out.append(row)
    return out


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def read_xlsx_rows(p: Path) -> tuple[list[str], list[dict]]:
    # XLSX 是 ZIP；仅限制压缩文件大小不足以防止 ZIP bomb。
    try:
        with zipfile.ZipFile(p) as archive:
            entries = archive.infolist()
            if len(entries) > MAX_XLSX_ENTRIES:
                raise KbSourceError("Excel 文件包含过多压缩条目")
            if sum(item.file_size for item in entries) > MAX_XLSX_EXPANDED_BYTES:
                raise KbSourceError("Excel 展开后超过 40MB 上限")
    except KbSourceError:
        raise
    except (zipfile.BadZipFile, OSError) as exc:
        raise KbSourceError("Excel 文件不是有效的 XLSX 压缩包") from exc

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise KbSourceError("服务器缺少 Excel 解析组件（openpyxl），请改用 CSV") from exc

    try:
        wb = load_workbook(io.BytesIO(p.read_bytes()), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 — 不把底层 XML/路径异常原样传给 UI
        raise KbSourceError("Excel 文件无法解析，请确认它是有效的 XLSX 文件") from exc

    try:
        if not wb.worksheets:
            raise KbSourceError("Excel 文件中没有工作表")
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            return [], []
        cols = [str(h).strip() if h is not None else f"列{i + 1}" for i, h in enumerate(header)]
        rows: list[dict] = []
        for row in rows_iter:
            if row is None or all(v is None or str(v).strip() == "" for v in row):
                continue
            rows.append({
                cols[i]: ("" if value is None else str(value))
                for i, value in enumerate(row) if i < len(cols)
            })
            if len(rows) > MAX_TOPIC_ROWS:
                raise KbSourceError(f"选题库超过 {MAX_TOPIC_ROWS:,} 行上限")
        return cols, rows
    except KbSourceError:
        raise
    except Exception as exc:  # noqa: BLE001 — 损坏的 XML 不应回传裸堆栈
        raise KbSourceError("Excel 文件内容损坏，无法读取") from exc
    finally:
        wb.close()
